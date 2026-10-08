"""Rendering a ReportDoc to PDF and to DOCX.

Both from the same structure, so the two cannot drift. Neither needs a
system library: reportlab and python-docx are pure-Python wheels, which
matters because a report generator that only works on the maintainer's
laptop is not a feature.

The DOCX is a real document — styled headings, native tables — rather than
a page screenshot. A client who asks for DOCX is going to edit it.
"""
from __future__ import annotations

import io
import re
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from .model import SEVERITY_COLOUR, Block, ReportDoc

AGENT_NOTICE = (
    "Parts of the narrative in this document were revised by a language "
    "model. Findings, severities and counts are taken directly from the "
    "engagement data and were not generated or altered by it."
)


def _esc(s: str | None) -> str:
    """Escape for reportlab's mini-markup, which treats < and & as tags."""
    if not s:
        return ""
    return (str(s).replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;"))


def _clean(s: str | None, limit: int = 4000) -> str:
    """Collapse scanner output into something that sets on a page.

    Plugin output arrives with hard-wrapped lines and runs of blank space
    that make a PDF three times longer than it needs to be.
    """
    if not s:
        return ""
    txt = re.sub(r"[ \t]+", " ", str(s))
    txt = re.sub(r"\n{3,}", "\n\n", txt).strip()
    if len(txt) > limit:
        txt = txt[:limit].rsplit(" ", 1)[0] + " […truncated]"
    return txt


# ------------------------------------------------------------------ PDF
def _styles():
    ss = getSampleStyleSheet()
    base = ss["BodyText"]
    base.fontName, base.fontSize, base.leading = "Helvetica", 9.5, 13.5
    base.alignment = TA_LEFT
    base.spaceAfter = 6
    return {
        "body": base,
        "h1": ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold",
                             fontSize=16, spaceBefore=4, spaceAfter=10,
                             textColor=colors.HexColor("#111827")),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold",
                             fontSize=11.5, spaceBefore=12, spaceAfter=4,
                             textColor=colors.HexColor("#1F2937")),
        "small": ParagraphStyle("small", parent=base, fontSize=8,
                                textColor=colors.HexColor("#6B7280")),
        "cell": ParagraphStyle("cell", parent=base, fontSize=8, leading=10.5,
                               spaceAfter=0),
        "cellhead": ParagraphStyle("cellhead", parent=base, fontSize=8,
                                   leading=10.5, spaceAfter=0,
                                   fontName="Helvetica-Bold",
                                   textColor=colors.white),
    }


def to_pdf(doc: ReportDoc) -> bytes:
    buf = io.BytesIO()
    st = _styles()
    page = A4
    margin = 18 * mm

    pdf = BaseDocTemplate(
        buf, pagesize=page, title=f"{doc.title} — {doc.project}",
        author=doc.generated_by, subject=doc.subtitle,
        leftMargin=margin, rightMargin=margin,
        topMargin=margin, bottomMargin=20 * mm)
    frame = Frame(margin, 20 * mm, page[0] - 2 * margin,
                  page[1] - margin - 24 * mm, id="body")

    def decorate(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#6B7280"))
        canvas.drawString(margin, 12 * mm, f"{doc.project} · {doc.title}")
        canvas.drawRightString(page[0] - margin, 12 * mm, f"Page {canvas.getPageNumber()}")
        canvas.setStrokeColor(colors.HexColor("#E5E7EB"))
        canvas.line(margin, 15 * mm, page[0] - margin, 15 * mm)
        canvas.restoreState()

    pdf.addPageTemplates([PageTemplate(id="main", frames=[frame], onPage=decorate)])

    flow: list = [
        Paragraph(_esc(doc.title), st["h1"]),
        Paragraph(_esc(doc.subtitle), st["body"]),
        Spacer(1, 4),
        Paragraph(
            f"{_esc(doc.project)} · generated "
            f"{doc.generated_at:%Y-%m-%d %H:%M} UTC by {_esc(doc.generated_by)}",
            st["small"]),
    ]
    if doc.agent_edited:
        flow += [Spacer(1, 6), Paragraph(_esc(AGENT_NOTICE), st["small"])]
    flow.append(Spacer(1, 10))

    for sec in doc.sections:
        if sec.page_break_before:
            flow.append(PageBreak())
        flow.append(Paragraph(_esc(sec.heading), st["h1"]))
        for b in sec.blocks:
            flow += _pdf_block(b, st, page[0] - 2 * margin)

    pdf.build(flow)
    return buf.getvalue()


def _pdf_block(b: Block, st, width: float) -> list:
    if b.kind == "para":
        return [Paragraph(_esc(_clean(b.text)).replace("\n", "<br/>"), st["body"])]

    if b.kind == "bullets":
        return [Paragraph(f"•&nbsp;&nbsp;{_esc(i)}", st["body"]) for i in b.items]

    if b.kind == "kv":
        t = Table([[Paragraph(_esc(k), st["cell"]), Paragraph(_esc(v), st["cell"])]
                   for k, v in b.pairs],
                  colWidths=[width * 0.45, width * 0.55], hAlign="LEFT")
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E5E7EB")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F9FAFB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return [t, Spacer(1, 8)]

    if b.kind == "table":
        n = len(b.headers) or 1
        # The first column is usually a hostname and needs the room.
        widths = ([width * 0.26] + [(width * 0.74) / (n - 1)] * (n - 1)
                  if n > 1 else [width])
        data = [[Paragraph(_esc(h), st["cellhead"]) for h in b.headers]]
        data += [[Paragraph(_esc(c), st["cell"]) for c in r] for r in b.rows]
        t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#374151")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E5E7EB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#F9FAFB")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return [t, Spacer(1, 8)]

    if b.kind == "findings":
        out: list = []
        for i, f in enumerate(b.findings, 1):
            colour = SEVERITY_COLOUR.get(f.severity, "#5A5A5A")
            where = f.host + (f":{f.port}" if f.port else "")
            head = Paragraph(
                f'<font color="{colour}"><b>{f.severity.upper()}</b></font>'
                f'&nbsp;&nbsp;{i}. {_esc(f.title)}', st["h2"])
            meta = Paragraph(
                f"<b>Host:</b> {_esc(where)}"
                + (f" &nbsp;·&nbsp; <b>Ref:</b> {_esc(f.external_id)}"
                   if f.external_id else ""), st["small"])
            bits = [head, meta]
            if f.description:
                bits.append(Paragraph(
                    _esc(_clean(f.description)).replace("\n", "<br/>"), st["body"]))
            # Always present, because "no remediation recorded" is itself
            # information the client needs — it means nobody wrote one.
            rem = _clean(f.remediation) or "No remediation was recorded for this finding."
            bits.append(Paragraph(
                f"<b>Recommended remediation:</b> "
                f"{_esc(rem).replace(chr(10), '<br/>')}", st["body"]))
            bits.append(Spacer(1, 6))
            # KeepTogether stops a finding splitting its heading from its
            # body across a page break.
            out.append(KeepTogether(bits))
        return out

    return []


# ----------------------------------------------------------------- DOCX
def to_docx(doc: ReportDoc) -> bytes:
    d = Document()

    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    d.add_heading(doc.title, level=0)
    p = d.add_paragraph(doc.subtitle)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    meta = d.add_paragraph()
    run = meta.add_run(
        f"{doc.project} · generated {doc.generated_at:%Y-%m-%d %H:%M} UTC "
        f"by {doc.generated_by}")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    if doc.agent_edited:
        n = d.add_paragraph()
        r = n.add_run(AGENT_NOTICE)
        r.italic = True
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    for sec in doc.sections:
        if sec.page_break_before:
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        d.add_heading(sec.heading, level=1)
        for b in sec.blocks:
            _docx_block(d, b)

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _docx_block(d: Document, b: Block) -> None:
    if b.kind == "para":
        d.add_paragraph(_clean(b.text))

    elif b.kind == "bullets":
        for i in b.items:
            d.add_paragraph(i, style="List Bullet")

    elif b.kind == "kv":
        t = d.add_table(rows=0, cols=2)
        t.style = "Light Grid Accent 1"
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        for k, v in b.pairs:
            cells = t.add_row().cells
            cells[0].text = k
            cells[1].text = v
            for para in cells[0].paragraphs:
                for run in para.runs:
                    run.bold = True
        d.add_paragraph()

    elif b.kind == "table":
        t = d.add_table(rows=1, cols=max(1, len(b.headers)))
        t.style = "Light Grid Accent 1"
        for i, h in enumerate(b.headers):
            cell = t.rows[0].cells[i]
            cell.text = h
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True
        for r in b.rows:
            cells = t.add_row().cells
            for i, c in enumerate(r[:len(cells)]):
                cells[i].text = str(c)
                for para in cells[i].paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(8)
        d.add_paragraph()

    elif b.kind == "findings":
        for i, f in enumerate(b.findings, 1):
            h = d.add_heading(level=2)
            sev = h.add_run(f"{f.severity.upper()}  ")
            hexcol = SEVERITY_COLOUR.get(f.severity, "#5A5A5A").lstrip("#")
            sev.font.color.rgb = RGBColor.from_string(hexcol)
            h.add_run(f"{i}. {f.title}")

            where = f.host + (f":{f.port}" if f.port else "")
            m = d.add_paragraph()
            mr = m.add_run("Host: " + where
                           + (f"   ·   Ref: {f.external_id}" if f.external_id else ""))
            mr.font.size = Pt(8)
            mr.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

            if f.description:
                d.add_paragraph(_clean(f.description))
            rp = d.add_paragraph()
            rp.add_run("Recommended remediation: ").bold = True
            rp.add_run(_clean(f.remediation)
                       or "No remediation was recorded for this finding.")


RENDERERS = {"pdf": to_pdf, "docx": to_docx}

MEDIA_TYPE = {
    "pdf": "application/pdf",
    "docx": ("application/vnd.openxmlformats-officedocument"
             ".wordprocessingml.document"),
}


def render(doc: ReportDoc, fmt: str) -> tuple[bytes, str]:
    fn = RENDERERS.get(fmt)
    if fn is None:
        raise ValueError(f"unknown format {fmt!r}; expected pdf or docx")
    return fn(doc), MEDIA_TYPE[fmt]


def filename(doc: ReportDoc, fmt: str, when: datetime | None = None) -> str:
    stamp = (when or doc.generated_at).strftime("%Y%m%d-%H%M")
    safe = re.sub(r"[^A-Za-z0-9._-]+", "-", f"{doc.project}-{doc.kind}").strip("-")
    return f"{safe}-{stamp}.{fmt}"
